from __future__ import annotations

import fcntl
import threading
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterator, Sequence

from openpyxl import Workbook, load_workbook

from app.core.diagnostics import diagnostic_operation
from app.core.domain import EntryPatch, EntryRef
from app.core.errors import StorageError, WorkbookConflictError, WorkbookRowNotFoundError
from app.core.models import VocabularyEntry, normalize_word_key
from app.core.search_index import build_entry_search_blob
from app.storage.excel_backup import backup_workbook
from app.storage.excel_reader import read_header
from app.storage.excel_schema import style_last_row, style_row, write_header
from app.storage.excel_serializer import (
    SHEET_NAME,
    is_managed_column,
    label_to_key,
    render_cell,
    row_to_entry,
)
from app.storage.excel_workbook_io import (
    load_for_write,
    raise_with_backup_hint,
    save_workbook,
)
from app.storage.workbook_revision import WorkbookRevision
from app.storage.workbook_snapshot import (
    CellState,
    RowSnapshot,
    SnapshotCell,
    SnapshotRow,
    WorkbookSnapshot,
    entry_ref,
    row_identity,
)


@dataclass(frozen=True)
class RepositoryDeleteOutcome:
    removed: int
    rows: tuple[RowSnapshot, ...]
    before_revision: WorkbookRevision
    after_revision: WorkbookRevision
    backup_path: Path | None = None


@dataclass(frozen=True)
class RepositoryRestoreOutcome:
    restored: int
    skipped_existing: int
    before_revision: WorkbookRevision
    after_revision: WorkbookRevision
    restored_row_numbers: tuple[int, ...] = ()


@dataclass(frozen=True)
class RepositoryPatchOutcome:
    snapshot: WorkbookSnapshot | None
    after_revision: WorkbookRevision
    backup_path: Path | None = None
    affected_refs: tuple[EntryRef, ...] = ()


class WorkbookRepository:
    def __init__(
        self,
        *,
        read_func: Callable | None = None,
        load_func: Callable = load_for_write,
        save_func: Callable = save_workbook,
        backup_func: Callable = backup_workbook,
    ) -> None:
        self._read = read_func or (
            lambda path: load_workbook(path, read_only=True, data_only=False)
        )
        self._load = load_func
        self._save = save_func
        self._backup = backup_func
        self._snapshot_lock = threading.RLock()
        self._snapshot_cache: dict[Path, WorkbookSnapshot] = {}

    def read_snapshot(self, path: Path) -> WorkbookSnapshot:
        with diagnostic_operation("workbook_read") as diagnostic:
            snapshot = self._read_snapshot(path)
            diagnostic.finish(
                "empty" if not snapshot.rows else "success",
                row_count=len(snapshot.rows),
            )
            return snapshot

    def _read_snapshot(self, path: Path) -> WorkbookSnapshot:
        path = _resolved_path(path)
        with self._snapshot_lock:
            current = WorkbookRevision.capture(path)
            cache_key = Path(current.resolved_path)
            cached = self._snapshot_cache.get(cache_key)
            if cached is not None and cached.revision == current:
                return cached
            for attempt in range(2):
                before = current if attempt == 0 else WorkbookRevision.capture(path)
                if not before.exists:
                    snapshot = WorkbookSnapshot(
                        path,
                        before,
                        SHEET_NAME,
                        (),
                        (),
                        (),
                    )
                    self._snapshot_cache[cache_key] = snapshot
                    return snapshot
                workbook = self._read(path)
                try:
                    if SHEET_NAME not in workbook.sheetnames:
                        snapshot = WorkbookSnapshot(
                            path,
                            before,
                            SHEET_NAME,
                            (),
                            (),
                            (),
                        )
                    else:
                        snapshot = self._snapshot_from_sheet(
                            path,
                            before,
                            workbook[SHEET_NAME],
                        )
                finally:
                    workbook.close()
                after = WorkbookRevision.capture(path)
                if after == before:
                    self._snapshot_cache[cache_key] = snapshot
                    return snapshot
                if attempt == 0:
                    current = after
                    continue
        raise WorkbookConflictError(f"Excel 파일이 읽는 동안 변경되었습니다: {path}")

    def invalidate_snapshot(self, path: Path) -> None:
        revision = WorkbookRevision.capture(_resolved_path(path))
        with self._snapshot_lock:
            self._snapshot_cache.pop(Path(revision.resolved_path), None)

    def find_entry(
        self,
        snapshot: WorkbookSnapshot,
        ref_or_key: EntryRef | tuple[str, str],
    ) -> SnapshotRow | None:
        if isinstance(ref_or_key, EntryRef):
            return snapshot.find_ref(ref_or_key)
        return snapshot.find_key(*ref_or_key)

    def commit_patch(
        self,
        path: Path,
        expected_revision: WorkbookRevision,
        patches: Sequence[EntryPatch],
        columns: Sequence[str],
    ) -> WorkbookSnapshot:
        return self._commit_patch(
            path,
            expected_revision,
            patches,
            columns,
            backup_reason=None,
            refresh_snapshot=True,
        ).snapshot

    def commit_patch_with_backup(
        self,
        path: Path,
        expected_revision: WorkbookRevision,
        patches: Sequence[EntryPatch],
        columns: Sequence[str],
        *,
        backup_reason: str,
        refresh_snapshot: bool = True,
    ) -> RepositoryPatchOutcome:
        return self._commit_patch(
            path,
            expected_revision,
            patches,
            columns,
            backup_reason=backup_reason,
            refresh_snapshot=refresh_snapshot,
        )

    def commit_patch_without_snapshot(
        self,
        path: Path,
        expected_revision: WorkbookRevision,
        patches: Sequence[EntryPatch],
        columns: Sequence[str],
        *,
        backup_reason: str | None = None,
    ) -> RepositoryPatchOutcome:
        return self._commit_patch(
            path,
            expected_revision,
            patches,
            columns,
            backup_reason=backup_reason,
            refresh_snapshot=False,
        )

    def _commit_patch(
        self,
        path: Path,
        expected_revision: WorkbookRevision,
        patches: Sequence[EntryPatch],
        columns: Sequence[str],
        *,
        backup_reason: str | None,
        refresh_snapshot: bool,
    ) -> RepositoryPatchOutcome:
        path = _resolved_path(path)
        if not patches:
            self._assert_revision(path, expected_revision)
            snapshot = self.read_snapshot(path) if refresh_snapshot else None
            return RepositoryPatchOutcome(
                snapshot,
                WorkbookRevision.capture(path),
            )
        backup_path: Path | None = None
        affected_row_numbers: list[int] = []
        with _workbook_lock(path):
            self._assert_revision(path, expected_revision)
            if backup_reason is not None and path.exists():
                backup_path = self._backup(path, backup_reason)
            workbook = self._load_or_create(path, columns)
            try:
                sheet = workbook[SHEET_NAME]
                file_columns = read_header(sheet) or list(columns)
                if not read_header(sheet):
                    write_header(sheet, file_columns)
                for patch in patches:
                    entry = patch.entry
                    if patch.ref is None:
                        sheet.append([render_cell(entry, key) for key in file_columns])
                        style_last_row(sheet, file_columns)
                        affected_row_numbers.append(sheet.max_row)
                        continue
                    row_number = self._resolve_ref(sheet, file_columns, patch.ref)
                    self._patch_managed_cells(
                        sheet,
                        row_number,
                        file_columns,
                        entry,
                    )
                    affected_row_numbers.append(row_number)
                affected_refs = tuple(
                    self._entry_ref_at_row(sheet, file_columns, row_number)
                    for row_number in affected_row_numbers
                )
                try:
                    self._persist(workbook, path, expected_revision, backup_path)
                except StorageError as exc:
                    if backup_path is not None:
                        raise_with_backup_hint(exc, backup_path)
                    raise
            finally:
                workbook.close()
        after_revision = WorkbookRevision.capture(path)
        self.invalidate_snapshot(path)
        snapshot = self.read_snapshot(path) if refresh_snapshot else None
        return RepositoryPatchOutcome(
            snapshot,
            after_revision,
            backup_path,
            affected_refs,
        )

    def delete_rows(
        self,
        path: Path,
        expected_revision: WorkbookRevision,
        refs: Sequence[EntryRef],
        *,
        create_backup: bool = True,
    ) -> RepositoryDeleteOutcome:
        path = _resolved_path(path)
        if not refs:
            current = WorkbookRevision.capture(path)
            self._assert_revision(path, expected_revision)
            return RepositoryDeleteOutcome(0, (), current, current)
        with _workbook_lock(path):
            self._assert_revision(path, expected_revision)
            workbook = self._load(path)
            backup_path: Path | None = None
            try:
                if SHEET_NAME not in workbook.sheetnames:
                    current = WorkbookRevision.capture(path)
                    return RepositoryDeleteOutcome(0, (), current, current)
                sheet = workbook[SHEET_NAME]
                columns = read_header(sheet)
                headers = self._headers(sheet)
                rows: list[RowSnapshot] = []
                resolved: set[int] = set()
                for ref in refs:
                    row_number = self._resolve_ref(sheet, columns, ref)
                    if row_number in resolved:
                        continue
                    resolved.add(row_number)
                    rows.append(self._capture_deleted_row(sheet, row_number, headers, columns))
                if not rows:
                    current = WorkbookRevision.capture(path)
                    return RepositoryDeleteOutcome(0, (), current, current)
                if create_backup:
                    backup_path = self._backup(path, "delete-rows")
                for row_number in sorted(resolved, reverse=True):
                    sheet.delete_rows(row_number, 1)
                try:
                    self._persist(workbook, path, expected_revision, backup_path)
                except StorageError as exc:
                    if backup_path is not None:
                        raise_with_backup_hint(exc, backup_path)
                    raise
            finally:
                workbook.close()
            after = WorkbookRevision.capture(path)
            self.invalidate_snapshot(path)
        return RepositoryDeleteOutcome(
            len(rows),
            tuple(sorted(rows, key=lambda row: row.original_row_number)),
            expected_revision,
            after,
            backup_path,
        )

    def delete_matching_rows(
        self,
        path: Path,
        expected_revision: WorkbookRevision,
        language: str,
        word_keys: set[str],
        *,
        create_backup: bool = True,
    ) -> RepositoryDeleteOutcome:
        """Compatibility key deletion without a redundant snapshot scan."""
        path = _resolved_path(path)
        if not word_keys:
            current = WorkbookRevision.capture(path)
            self._assert_revision(path, expected_revision)
            return RepositoryDeleteOutcome(0, (), current, current)
        with _workbook_lock(path):
            self._assert_revision(path, expected_revision)
            if not path.exists():
                current = WorkbookRevision.capture(path)
                return RepositoryDeleteOutcome(0, (), current, current)
            workbook = self._load(path)
            backup_path: Path | None = None
            try:
                if SHEET_NAME not in workbook.sheetnames:
                    current = WorkbookRevision.capture(path)
                    return RepositoryDeleteOutcome(0, (), current, current)
                sheet = workbook[SHEET_NAME]
                columns = read_header(sheet)
                if "language" not in columns or "word" not in columns:
                    current = WorkbookRevision.capture(path)
                    return RepositoryDeleteOutcome(0, (), current, current)
                headers = self._headers(sheet)
                language_column = columns.index("language") + 1
                word_column = columns.index("word") + 1
                rows: list[RowSnapshot] = []
                for row_number in range(2, sheet.max_row + 1):
                    row_language = sheet.cell(
                        row=row_number,
                        column=language_column,
                    ).value
                    row_word = sheet.cell(row=row_number, column=word_column).value
                    if row_language != language or not isinstance(row_word, str):
                        continue
                    key = normalize_word_key(row_word, language)  # type: ignore[arg-type]
                    if key in word_keys:
                        rows.append(
                            self._capture_deleted_row(
                                sheet,
                                row_number,
                                headers,
                                columns,
                            )
                        )
                if not rows:
                    current = WorkbookRevision.capture(path)
                    return RepositoryDeleteOutcome(0, (), current, current)
                if create_backup:
                    backup_path = self._backup(path, "delete-rows")
                for row in reversed(rows):
                    sheet.delete_rows(row.original_row_number, 1)
                try:
                    self._persist(workbook, path, expected_revision, backup_path)
                except StorageError as exc:
                    if backup_path is not None:
                        raise_with_backup_hint(exc, backup_path)
                    raise
            finally:
                workbook.close()
            after = WorkbookRevision.capture(path)
            self.invalidate_snapshot(path)
        return RepositoryDeleteOutcome(
            len(rows),
            tuple(rows),
            expected_revision,
            after,
            backup_path,
        )

    def restore_rows(
        self,
        path: Path,
        expected_revision: WorkbookRevision,
        rows: Sequence[RowSnapshot],
    ) -> RepositoryRestoreOutcome:
        path = _resolved_path(path)
        if not rows:
            current = WorkbookRevision.capture(path)
            self._assert_revision(path, expected_revision)
            return RepositoryRestoreOutcome(0, 0, current, current)
        with _workbook_lock(path):
            self._assert_revision(path, expected_revision)
            workbook = self._load(path)
            restored = 0
            skipped = 0
            restored_row_numbers: list[int] = []
            try:
                if SHEET_NAME not in workbook.sheetnames:
                    sheet = workbook.create_sheet(SHEET_NAME)
                    labels = [cell.header for cell in rows[0].cells]
                    sheet.append(labels)
                else:
                    sheet = workbook[SHEET_NAME]
                headers = self._headers(sheet)
                for row in sorted(rows, key=lambda item: item.original_row_number):
                    if self._has_equivalent_row(sheet, headers, row):
                        skipped += 1
                        continue
                    target = min(max(2, row.original_row_number), sheet.max_row + 1)
                    sheet.insert_rows(target, 1)
                    self._restore_row(sheet, target, headers, row)
                    restored += 1
                    restored_row_numbers.append(row.original_row_number)
                if restored:
                    self._persist(workbook, path, expected_revision, None)
            finally:
                workbook.close()
            after = WorkbookRevision.capture(path)
            self.invalidate_snapshot(path)
        return RepositoryRestoreOutcome(
            restored,
            skipped,
            expected_revision,
            after,
            tuple(restored_row_numbers),
        )

    def _snapshot_from_sheet(
        self,
        path: Path,
        revision: WorkbookRevision,
        sheet,
    ) -> WorkbookSnapshot:
        rows_iterator = sheet.iter_rows(values_only=False)
        try:
            header_cells = next(rows_iterator)
        except StopIteration:
            return WorkbookSnapshot(path, revision, SHEET_NAME, (), (), ())
        headers = tuple(cell.value for cell in header_cells)
        columns = tuple(label_to_key(header) for header in headers)
        rows: list[SnapshotRow] = []
        for row_number, source_cells in enumerate(rows_iterator, start=2):
            cells = tuple(
                SnapshotCell(
                    header,
                    columns[index - 1],
                    source_cells[index - 1].value,
                    source_cells[index - 1].data_type,
                )
                for index, header in enumerate(headers, start=1)
            )
            values = tuple(cell.value for cell in cells)
            if not any(value is not None for value in values):
                continue
            entry = row_to_entry(list(columns), values)
            ref = entry_ref(
                entry,
                row_number,
                values,
                tuple(cell.data_type for cell in cells),
            )
            rows.append(
                SnapshotRow(
                    ref,
                    entry.to_json(),
                    cells,
                    build_entry_search_blob(entry),
                )
            )
        return WorkbookSnapshot(
            path,
            revision,
            SHEET_NAME,
            headers,
            columns,
            tuple(rows),
        )

    def _capture_deleted_row(
        self,
        sheet,
        row_number: int,
        headers: tuple[object, ...],
        columns: list[str],
    ) -> RowSnapshot:
        cells = tuple(
            CellState.capture(
                sheet.cell(row=row_number, column=index),
                header,
                columns[index - 1],
            )
            for index, header in enumerate(headers, start=1)
        )
        values = tuple(cell.value for cell in cells)
        entry = row_to_entry(columns, values)
        ref = entry_ref(
            entry,
            row_number,
            values,
            tuple(cell.data_type for cell in cells),
        )
        dimension = sheet.row_dimensions[row_number]
        return RowSnapshot(
            ref=ref,
            original_row_number=row_number,
            entry_json=entry.to_json(),
            cells=cells,
            height=dimension.height,
            hidden=bool(dimension.hidden),
            outline_level=int(dimension.outlineLevel or 0),
        )

    def _resolve_ref(self, sheet, columns: list[str], ref: EntryRef) -> int:
        row_number = ref.row_number
        if row_number < 2 or row_number > sheet.max_row:
            raise WorkbookRowNotFoundError(f"Excel 행을 찾을 수 없습니다: {ref.word_key}")
        values = tuple(
            sheet.cell(row=row_number, column=index).value for index in range(1, len(columns) + 1)
        )
        data_types = tuple(
            sheet.cell(row=row_number, column=index).data_type
            for index in range(1, len(columns) + 1)
        )
        if row_identity(values, data_types) != ref.identity:
            raise WorkbookRowNotFoundError(f"Excel 행이 변경되었습니다: {ref.word_key}")
        return row_number

    def _patch_managed_cells(
        self,
        sheet,
        row_number: int,
        columns: Sequence[str],
        entry: VocabularyEntry,
    ) -> None:
        for column_index, key in enumerate(columns, start=1):
            if is_managed_column(key):
                sheet.cell(
                    row=row_number,
                    column=column_index,
                    value=render_cell(entry, key),
                )
        style_row(sheet, row_number, list(columns))

    @staticmethod
    def _entry_ref_at_row(sheet, columns: Sequence[str], row_number: int) -> EntryRef:
        cells = tuple(
            sheet.cell(row=row_number, column=index) for index in range(1, len(columns) + 1)
        )
        # openpyxl serializes explicit empty strings as empty inline strings;
        # a subsequent read therefore yields ``None``/``inlineStr``. Normalize
        # here so the identity returned by commit matches the saved workbook.
        values = tuple(None if cell.value == "" else cell.value for cell in cells)
        data_types = tuple(
            "inlineStr"
            if cell.value == "" and cell.data_type in {"s", "inlineStr"}
            else cell.data_type
            for cell in cells
        )
        return entry_ref(
            row_to_entry(list(columns), values),
            row_number,
            values,
            data_types,
        )

    def _has_equivalent_row(
        self,
        sheet,
        headers: tuple[object, ...],
        snapshot: RowSnapshot,
    ) -> bool:
        expected = {str(cell.header): cell.value for cell in snapshot.cells}
        for row_number in range(2, sheet.max_row + 1):
            if all(
                sheet.cell(row=row_number, column=index).value == expected.get(str(header))
                for index, header in enumerate(headers, start=1)
                if str(header) in expected
            ):
                return True
        return False

    def _restore_row(
        self,
        sheet,
        row_number: int,
        headers: tuple[object, ...],
        snapshot: RowSnapshot,
    ) -> None:
        by_label = {str(cell.header): cell for cell in snapshot.cells}
        by_key = {cell.key: cell for cell in snapshot.cells if cell.key}
        for column_index, header in enumerate(headers, start=1):
            state = by_label.get(str(header)) or by_key.get(label_to_key(header))
            if state is not None:
                state.apply(sheet.cell(row=row_number, column=column_index))
        dimension = sheet.row_dimensions[row_number]
        dimension.height = snapshot.height
        dimension.hidden = snapshot.hidden
        dimension.outlineLevel = snapshot.outline_level

    @staticmethod
    def _headers(sheet) -> tuple[object, ...]:
        return tuple(
            sheet.cell(row=1, column=index).value for index in range(1, sheet.max_column + 1)
        )

    @staticmethod
    def _assert_revision(path: Path, expected: WorkbookRevision) -> None:
        current = WorkbookRevision.capture(path)
        if current != expected:
            raise WorkbookConflictError(f"Excel 파일이 외부에서 변경되었습니다: {path}")

    def _persist(
        self,
        workbook,
        path: Path,
        before_revision: WorkbookRevision,
        backup_path: Path | None,
    ) -> None:
        if self._save is save_workbook:
            self._save(
                workbook,
                path,
                before_revision=before_revision,
                backup_path=backup_path,
            )
            return
        self._save(workbook, path)

    def _load_or_create(self, path: Path, columns: Sequence[str]):
        if path.exists():
            workbook = self._load(path)
            if SHEET_NAME not in workbook.sheetnames:
                sheet = workbook.create_sheet(SHEET_NAME)
                write_header(sheet, list(columns))
            return workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        write_header(sheet, list(columns))
        return workbook


@contextmanager
def _workbook_lock(path: Path) -> Iterator[None]:
    path.parent.mkdir(parents=True, exist_ok=True)
    lock_path = path.with_name(f".{path.name}.jelly.lock")
    with lock_path.open("a+b") as lock_file:
        fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX)
        try:
            yield
        finally:
            fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)


def _resolved_path(path: Path) -> Path:
    return Path(path).expanduser().resolve(strict=False)
