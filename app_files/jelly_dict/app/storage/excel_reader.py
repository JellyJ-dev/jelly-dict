"""Read-side helpers for the vocabulary Excel file.

These functions never mutate the workbook on disk. For mutations see
`app.storage.excel_writer`.
"""

from __future__ import annotations

from dataclasses import dataclass
from enum import Enum
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.errors import ExcelFormatError, ExcelLockedError, StorageError
from app.core.models import VocabularyEntry, normalize_word_key
from app.storage.excel_serializer import (
    SHEET_NAME,
    label_to_key,
    row_to_entry,
)


class WorkbookReadState(str, Enum):
    OK = "ok"
    MISSING = "missing"
    ERROR = "error"


@dataclass(frozen=True)
class WorkbookReadResult:
    state: WorkbookReadState
    entries: tuple[VocabularyEntry, ...] = ()
    error: StorageError | None = None

    @property
    def ok(self) -> bool:
        return self.state is WorkbookReadState.OK


def list_entries(path: Path) -> list[VocabularyEntry]:
    """Compatibility reader: missing/unreadable files appear empty."""
    result = read_entries_result(path)
    return list(result.entries) if result.ok else []


def list_entries_strict(path: Path) -> list[VocabularyEntry]:
    """Return entries, raising a typed storage error for unreadable files."""
    result = read_entries_result(path)
    if result.error is not None:
        raise result.error
    return list(result.entries)


def read_entries_result(path: Path) -> WorkbookReadResult:
    """Read a workbook without conflating a missing file with read failure."""
    if not path.exists():
        return WorkbookReadResult(WorkbookReadState.MISSING)
    try:
        wb = load_workbook(path, read_only=True)
    except Exception as exc:
        return WorkbookReadResult(
            WorkbookReadState.ERROR,
            error=_typed_read_error(path, exc),
        )
    try:
        if SHEET_NAME not in wb.sheetnames:
            return WorkbookReadResult(WorkbookReadState.OK)
        ws = wb[SHEET_NAME]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return WorkbookReadResult(WorkbookReadState.OK)
        keys = [label_to_key(label) for label in header]
        out: list[VocabularyEntry] = []
        for raw in rows:
            if raw is None or not any(c is not None for c in raw):
                continue
            out.append(row_to_entry(keys, raw))
        return WorkbookReadResult(WorkbookReadState.OK, tuple(out))
    except Exception as exc:
        return WorkbookReadResult(
            WorkbookReadState.ERROR,
            error=_typed_read_error(path, exc),
        )
    finally:
        wb.close()


def find_existing(path: Path, language: str, word_key: str) -> VocabularyEntry | None:
    """Return a thin VocabularyEntry from the row matching (language, word_key).

    The Excel sheet is the source of truth for data the user may have
    edited by hand, so we read the row instead of the cache.
    """
    if not path.exists():
        return None
    try:
        wb = load_workbook(path, read_only=True)
    except Exception:
        return None
    try:
        if SHEET_NAME not in wb.sheetnames:
            return None
        rows = wb[SHEET_NAME].iter_rows(values_only=True)
        try:
            keys = [label_to_key(label) for label in next(rows)]
        except StopIteration:
            return None
        if "language" not in keys or "word" not in keys:
            return None
        lang_idx = keys.index("language")
        word_idx = keys.index("word")
        for raw in rows:
            if raw is None or len(raw) <= max(lang_idx, word_idx):
                continue
            row_language = raw[lang_idx]
            row_word = raw[word_idx]
            if row_language != language or not isinstance(row_word, str):
                continue
            if normalize_word_key(row_word, language) == word_key:  # type: ignore[arg-type]
                return row_to_entry(keys, raw)
        return None
    except Exception:
        return None
    finally:
        wb.close()


def _typed_read_error(path: Path, exc: Exception) -> StorageError:
    message = f"Excel 읽기 실패 ({path.name}): {exc}"
    if isinstance(exc, PermissionError):
        return ExcelLockedError(message)
    if isinstance(exc, (BadZipFile, InvalidFileException, KeyError, ValueError)):
        return ExcelFormatError(message)
    return StorageError(message)


def read_header(ws) -> list[str]:
    """Return the canonical field-key list inferred from row 1."""
    if ws.max_row < 1:
        return []
    labels = [ws.cell(row=1, column=idx).value for idx in range(1, ws.max_column + 1)]
    if not any(labels):
        return []
    return [label_to_key(label) for label in labels]
