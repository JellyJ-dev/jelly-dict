#!/usr/bin/env python3
"""Read-only runtime checks used by shell entrypoints.

The checker never installs packages, writes configuration, or starts the Jelly Dict
GUI.  Keeping probes here makes quickstart orchestration testable without embedding
Python programs in the shell script.
"""

from __future__ import annotations

import importlib.util
import os
import re
import sys
from importlib import metadata
from io import BytesIO
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
APP_DIR = SCRIPT_DIR.parent / "jelly_dict"
sys.path.insert(0, str(APP_DIR))

from app.core.dependency_policy import (  # noqa: E402
    CONSTRAINT_PINS,
    KOKORO_REQUIRED_FILES,
    KOKORO_REQUIRED_MODULES,
    PYTHON_MAX_EXCLUSIVE,
    PYTHON_MIN,
    REQUIREMENT_GROUPS,
    RequirementRule,
    constraint_path,
)


def _version_tuple(value: str) -> tuple[int, ...]:
    return tuple(int(part) for part in re.findall(r"\d+", value)[:3])


def _active_rules(group: str) -> tuple[RequirementRule, ...]:
    return tuple(
        rule for rule in REQUIREMENT_GROUPS[group] if not rule.marker or rule.marker == sys.platform
    )


def _compare(left: tuple[int, ...], right: tuple[int, ...]) -> int:
    width = max(len(left), len(right))
    padded_left = left + (0,) * (width - len(left))
    padded_right = right + (0,) * (width - len(right))
    return (padded_left > padded_right) - (padded_left < padded_right)


def _satisfies(installed: str, specifier: str) -> bool:
    value = _version_tuple(installed)
    for clause in specifier.split(","):
        match = re.fullmatch(r"(>=|<=|==|!=|>|<)\s*([0-9][0-9.]*)", clause.strip())
        if match is None:
            raise ValueError(f"unsupported dependency specifier: {clause}")
        relation = _compare(value, _version_tuple(match.group(2)))
        operator = match.group(1)
        if operator == ">=" and relation < 0:
            return False
        if operator == "<=" and relation > 0:
            return False
        if operator == ">" and relation <= 0:
            return False
        if operator == "<" and relation >= 0:
            return False
        if operator == "==" and relation != 0:
            return False
        if operator == "!=" and relation == 0:
            return False
    return True


def _check_distributions(group: str) -> list[str]:
    bad: list[str] = []
    for rule in _active_rules(group):
        if importlib.util.find_spec(rule.module) is None:
            bad.append(f"{rule.module} missing")
            continue
        try:
            installed = metadata.version(rule.distribution)
        except metadata.PackageNotFoundError:
            bad.append(f"{rule.distribution} metadata missing")
            continue
        if not _satisfies(installed, rule.specifier):
            bad.append(f"{rule.distribution} {installed} does not satisfy {rule.specifier}")
    return bad


def _constraint_pin_errors(group: str) -> list[str]:
    python_minor = f"{sys.version_info.major}.{sys.version_info.minor}"
    pins = CONSTRAINT_PINS.get(python_minor, {}).get(group, ())
    bad: list[str] = []
    for pin in pins:
        requirement, _, marker = pin.partition(";")
        if marker and "Darwin" in marker and sys.platform != "darwin":
            continue
        distribution, expected = requirement.strip().split("==", 1)
        try:
            installed = metadata.version(distribution)
        except metadata.PackageNotFoundError:
            bad.append(f"{distribution} metadata missing")
            continue
        if installed != expected:
            bad.append(f"{distribution} {installed} does not match constraint {expected}")
    return bad


def check_python_version() -> int:
    current = sys.version_info[:2]
    if current < PYTHON_MIN:
        print(
            f"  ✗ python3 >= {PYTHON_MIN[0]}.{PYTHON_MIN[1]} required, "
            f"found {sys.version.split()[0]}"
        )
        return 1
    if current >= PYTHON_MAX_EXCLUSIVE:
        print(f"  ✗ Python {current[0]}.{current[1]} is not used for Jelly Dict.app yet.")
        print("    Python 3.13 + Qt can abort while creating the macOS Cocoa platform plugin")
        print("    on some macOS 26 systems. Install Python 3.12 or 3.11 and rerun installer.")
        return 1
    print(f"  ✓ python3 version OK: {sys.version.split()[0]}")
    return 0


def python_supported() -> int:
    current = sys.version_info[:2]
    return 0 if PYTHON_MIN <= current < PYTHON_MAX_EXCLUSIVE else 1


def python_version_text() -> int:
    print(sys.version.split()[0])
    return 0


def check_packages() -> int:
    modules = tuple(rule.module for rule in _active_rules("runtime"))
    missing = [name for name in modules if importlib.util.find_spec(name) is None]
    if missing:
        print("  ✗ missing Python packages: " + ", ".join(missing))
        return 1

    bad = _check_distributions("runtime") + _constraint_pin_errors("runtime")
    if bad:
        print("  ✗ Python package version mismatch:")
        for item in bad:
            print("    - " + item)
        return 1
    print("  ✓ required Python packages and versions")
    return 0


def check_qt() -> int:
    spec = importlib.util.find_spec("PySide6")
    if spec is None or spec.origin is None:
        print("  ✗ PySide6 missing")
        return 1

    root = Path(spec.origin).resolve().parent
    plugins = root / "Qt" / "plugins"
    platforms = plugins / "platforms"
    cocoa = platforms / "libqcocoa.dylib"
    if not cocoa.exists():
        print(f"  ✗ Qt Cocoa platform plugin missing: {cocoa}")
        return 1

    os.environ["QT_PLUGIN_PATH"] = str(plugins)
    os.environ["QT_QPA_PLATFORM_PLUGIN_PATH"] = str(platforms)
    os.environ.setdefault("QT_MAC_WANTS_LAYER", "1")

    from PySide6 import QtCore

    print(f"  ✓ Qt runtime: {QtCore.qVersion()} / cocoa plugin")
    return 0


def check_playwright_chromium() -> int:
    if importlib.util.find_spec("playwright") is None:
        return 1

    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        path = Path(playwright.chromium.executable_path)
        if not path.exists():
            return 1
    print("  ✓ Playwright Chromium")
    return 0


def check_tts_packages() -> int:
    required = KOKORO_REQUIRED_MODULES
    missing = [name for name in required if importlib.util.find_spec(name) is None]
    if missing:
        print("  ✗ missing TTS Python packages: " + ", ".join(missing))
        return 1

    bad = _check_distributions("tts") + _constraint_pin_errors("tts")
    if bad:
        print("  ✗ TTS package version mismatch:")
        for item in bad:
            print("    - " + item)
        return 1

    root = Path.home() / ".cache" / "huggingface" / "hub" / "models--hexgrad--Kokoro-82M"
    ref = root / "refs" / "main"
    snapshot = (
        root / "snapshots" / ref.read_text(encoding="utf-8").strip() if ref.exists() else None
    )
    if snapshot is None or not snapshot.exists():
        missing_files = list(KOKORO_REQUIRED_FILES)
    else:
        missing_files = [name for name in KOKORO_REQUIRED_FILES if not (snapshot / name).exists()]
    if missing_files:
        print("  ✗ missing Kokoro local model cache: " + ", ".join(missing_files))
        return 1
    print("  ✓ TTS Python packages and versions")
    return 0


def print_constraint_path(group: str) -> int:
    python_minor = f"{sys.version_info.major}.{sys.version_info.minor}"
    try:
        path = constraint_path(group, python_minor)
    except KeyError:
        print(
            f"no {group} constraints for Python {python_minor}",
            file=sys.stderr,
        )
        return 1
    if not path.is_file():
        print(f"constraint file missing: {path}", file=sys.stderr)
        return 1
    print(path)
    return 0


def print_pyside_root() -> int:
    spec = importlib.util.find_spec("PySide6")
    if spec is None or spec.origin is None:
        return 1
    print(Path(spec.origin).resolve().parent)
    return 0


def dependency_smoke() -> int:
    from bs4 import BeautifulSoup
    from lxml import etree
    from openpyxl import Workbook, load_workbook
    from playwright.sync_api import sync_playwright
    from PySide6 import QtCore

    parsed = BeautifulSoup("<p data-jelly='1'>ok</p>", "lxml")
    if parsed.p is None or parsed.p.get_text() != "ok":
        raise RuntimeError("BeautifulSoup/lxml parser smoke failed")
    if etree.fromstring(b"<root><item/></root>").find("item") is None:
        raise RuntimeError("lxml etree smoke failed")

    stream = BytesIO()
    workbook = Workbook()
    workbook.active["A1"] = "jelly"
    workbook.save(stream)
    stream.seek(0)
    reopened = load_workbook(stream, read_only=True)
    try:
        if reopened.active["A1"].value != "jelly":
            raise RuntimeError("openpyxl in-memory round-trip failed")
    finally:
        reopened.close()

    if not QtCore.qVersion() or not callable(sync_playwright):
        raise RuntimeError("Qt/Playwright API smoke failed")
    print("  ✓ parser/browser/Qt dependency smoke")
    return 0


def main(argv: list[str] | None = None) -> int:
    args = sys.argv[1:] if argv is None else argv
    if not args:
        print(f"usage: {Path(sys.argv[0]).name} <command>", file=sys.stderr)
        return 2
    command, *command_args = args
    no_arg_commands = {
        "python-version": check_python_version,
        "python-supported": python_supported,
        "python-version-text": python_version_text,
        "packages": check_packages,
        "qt": check_qt,
        "pyside-root": print_pyside_root,
        "playwright-chromium": check_playwright_chromium,
        "tts-packages": check_tts_packages,
        "dependency-smoke": dependency_smoke,
    }
    if command in no_arg_commands and not command_args:
        return no_arg_commands[command]()
    if command == "constraint-path" and len(command_args) == 1:
        if command_args[0] not in REQUIREMENT_GROUPS:
            return 2
        return print_constraint_path(command_args[0])
    print(f"unknown checker command or arguments: {' '.join(args)}", file=sys.stderr)
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
